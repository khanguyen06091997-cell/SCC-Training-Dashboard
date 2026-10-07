import sys,zipfile,re,io,uuid
from pptx import Presentation
from pptx.util import Emu
from lxml import etree
def imgs(path):
    z=zipfile.ZipFile(path); order=[]
    pres=z.read('ppt/presentation.xml').decode(); rels=z.read('ppt/_rels/presentation.xml.rels').decode()
    rid2t=dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"',rels))|{a:b for b,a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"',rels)}
    for rid in re.findall(r'<p:sldId [^>]*r:id="(rId\d+)"',pres):
        sl='ppt/'+rid2t[rid]; srel=sl.replace('slides/','slides/_rels/')+'.rels'
        t=[x for x in re.findall(r'Target="([^"]+)"',z.read(srel).decode()) if 'media' in x][0]
        order.append(z.read('ppt/slides/'+t.replace('../','../').split('../')[-1] if False else 'ppt/'+t.replace('../','')))
    return order
tr,kp,out=sys.argv[1:4]
secs=[('1. Training',imgs(tr)),('2. KPI Setting',imgs(kp))]
p=Presentation(); p.slide_width=Emu(12192000); p.slide_height=Emu(6858000)
blank=p.slide_layouts[6]; ids=[]
for name,lst in secs:
    cur=[]
    for b in lst:
        s=p.slides.add_slide(blank); s.shapes.add_picture(io.BytesIO(b),0,0,p.slide_width,p.slide_height)
        cur.append(s.slide_id)
    ids.append((name,cur))
NS='http://schemas.openxmlformats.org/presentationml/2006/main'; P14='http://schemas.microsoft.com/office/powerpoint/2010/main'
root=p.part._element
ext=etree.SubElement(etree.SubElement(root,'{%s}extLst'%NS) if root.find('{%s}extLst'%NS) is None else root.find('{%s}extLst'%NS),'{%s}ext'%NS,uri='{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
sl=etree.SubElement(ext,'{%s}sectionLst'%P14,nsmap={'p14':P14})
for name,cur in ids:
    se=etree.SubElement(sl,'{%s}section'%P14,name=name,id='{%s}'%str(uuid.uuid4()).upper())
    li=etree.SubElement(se,'{%s}sldIdLst'%P14)
    for i in cur: etree.SubElement(li,'{%s}sldId'%P14,id=str(i))
p.core_properties.title='SCC Training Dashboard & KPI'; p.core_properties.author='Phòng Nhân sự · L&OD'
p.save(out); print([ (n,len(c)) for n,c in ids])
